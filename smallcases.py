from datetime import datetime
import os
import json
import math
import uuid
import datetime as dt
from functools import wraps
from flask import Blueprint, render_template, request, jsonify, session

smallcases_bp = Blueprint('smallcases', __name__)

_FETCHER = None
_LOGIN_REQUIRED = None
_DATA_DIR = '.'


def init_smallcases(fetcher, login_required, data_dir):
    global _FETCHER, _LOGIN_REQUIRED, _DATA_DIR
    _FETCHER = fetcher
    _LOGIN_REQUIRED = login_required
    _DATA_DIR = data_dir or '.'


def _now():
    return dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')


def _today():
    return dt.date.today().strftime('%Y-%m-%d')


def _safe_user(username=None):
    username = username or session.get('username')
    return ''.join(c for c in str(username or '') if c.isalnum() or c in ('_', '-')).lower()


def _file(username=None):
    return os.path.join(_DATA_DIR, 'user_data', f'{_safe_user(username)}_smallcases.json')


def load_user_smallcases(username=None):
    user = username or session.get('username')
    path = _file(username)
    from db_store import load_user_doc_db
    data = load_user_doc_db("smallcases", user, path, default_factory=dict)
    return data if isinstance(data, dict) else {}


def save_user_smallcases(data, username=None):
    user = username or session.get('username')
    filepath = _file(username)
    from db_store import save_user_doc_db
    save_user_doc_db("smallcases", user, data, filepath)


def _load():
    return load_user_smallcases()


def _save(data):
    save_user_smallcases(data)


def _id(prefix):
    return f'{prefix}_{uuid.uuid4().hex[:14]}'


def _num(value, field, allow_zero=False):
    try:
        n = float(value)
    except (TypeError, ValueError):
        raise ValueError(f'{field} must be a number')
    if not math.isfinite(n):
        raise ValueError(f'{field} must be a finite number')
    if allow_zero:
        if n < 0:
            raise ValueError(f'{field} cannot be negative')
    elif n <= 0:
        raise ValueError(f'{field} must be greater than zero')
    return n


def _date(value, field):
    if not value:
        raise ValueError(f'{field} is required')
    try:
        return dt.datetime.strptime(str(value), '%Y-%m-%d').date()
    except ValueError:
        raise ValueError(f'{field} must be YYYY-MM-DD')


def _symbol(value):
    s = str(value or '').strip().upper()
    if not s:
        raise ValueError('Stock symbol is required')
    # NSE and BSE symbols for the same stock are treated as one holding.
    # Keep .NS as the canonical market symbol so RELIANCE.NS and RELIANCE.BO
    # are merged for holdings, validation, duplicate detection and valuation.
    if s.endswith('.NS') or s.endswith('.BO'):
        s = s.rsplit('.', 1)[0] + '.NS'
    else:
        s += '.NS'
    return s


def _import_date(value):
    if value is None or str(value).strip() == '':
        raise ValueError('Date is required')
    if isinstance(value, (dt.datetime, dt.date)):
        return value.date() if isinstance(value, dt.datetime) else value
    raw = str(value).strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d'):
        try:
            return dt.datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    raise ValueError('Date must be YYYY-MM-DD or DD/MM/YYYY')


def _import_fingerprint(tx):
    return (
        _symbol(tx.get('symbol')),
        str(tx.get('transaction_date') or ''),
        str(tx.get('transaction_type') or '').upper(),
        round(float(tx.get('quantity', 0) or 0), 6),
        round(float(tx.get('price', 0) or 0), 6),
    )


def _case_key(data, identifier):
    if identifier in data:
        return identifier, data[identifier]
    for key, sc in data.items():
        if sc.get('id') == identifier:
            return key, sc
    return None, None


def _transactions(sc):
    return sc.setdefault('transactions', [])


def _audit(sc, action, entity_type='', entity_id='', old=None, new=None, details=''):
    sc.setdefault('audit_logs', []).insert(0, {
        'id': _id('audit'),
        'action': action,
        'entity_type': entity_type,
        'entity_id': entity_id,
        'old_value': old,
        'new_value': new,
        'details': details,
        'created_at': _now(),
        'created_by': session.get('username')
    })


def _ordered_transactions(sc):
    return sorted(enumerate(_transactions(sc)), key=lambda x: (x[1].get('transaction_date', ''), x[1].get('created_at', ''), x[0]))


def _position_state(sc):
    """Return FIFO lots and per-symbol aggregate state without changing history."""
    state = {}
    for _, tx in _ordered_transactions(sc):
        sym = _symbol(tx.get('symbol'))
        state.setdefault(sym, {'name': tx.get('stock_name') or sym, 'lots': [], 'buys': 0.0, 'buy_cost': 0.0, 'sells': 0.0, 'sell_proceeds': 0.0, 'realized_pnl': 0.0})
        st = state[sym]
        qty = float(tx.get('quantity', 0) or 0)
        price = float(tx.get('price', 0) or 0)
        charges = float(tx.get('charges', 0) or 0)
        if tx.get('transaction_type') == 'BUY':
            st['buys'] += qty
            st['buy_cost'] += qty * price + charges
            st['lots'].append({'qty': qty, 'price': price, 'charges_per_share': charges / qty if qty else 0.0, 'date': tx.get('transaction_date')})
        elif tx.get('transaction_type') == 'SELL':
            proceeds = qty * price - charges
            st['sells'] += qty
            st['sell_proceeds'] += proceeds
            remaining = qty
            cost = 0.0
            while remaining > 1e-9 and st['lots']:
                lot = st['lots'][0]
                take = min(remaining, lot['qty'])
                cost += take * (lot['price'] + lot['charges_per_share'])
                lot['qty'] -= take
                remaining -= take
                if lot['qty'] <= 1e-9:
                    st['lots'].pop(0)
            st['realized_pnl'] += proceeds - cost
    for st in state.values():
        st['remaining_qty'] = sum(max(0.0, x['qty']) for x in st['lots'])
        st['invested_remaining'] = sum(x['qty'] * (x['price'] + x['charges_per_share']) for x in st['lots'])
        st['avg_buy'] = st['invested_remaining'] / st['remaining_qty'] if st['remaining_qty'] else 0.0
    return state


def _validate_transactions(sc, candidate=None):
    txs = candidate if candidate is not None else _transactions(sc)
    # Validate chronological inventory and date ordering. Historical records are never rewritten.
    lots = {}
    last_date = {}
    for idx, tx in sorted(enumerate(txs), key=lambda x: (x[1].get('transaction_date', ''), x[1].get('created_at', ''), x[0])):
        typ = str(tx.get('transaction_type', '')).upper()
        sym = _symbol(tx.get('symbol'))
        qty = _num(tx.get('quantity'), 'quantity')
        price = _num(tx.get('price'), 'price')
        d = _date(tx.get('transaction_date'), 'transaction date')
        charges = _num(tx.get('charges', 0), 'charges', allow_zero=True)
        if typ not in ('BUY', 'SELL'):
            raise ValueError('transaction_type must be BUY or SELL')
        if typ == 'SELL':
            available = sum(x for x in lots.get(sym, []))
            if qty > available + 1e-9:
                raise ValueError(f'Cannot sell {qty:g} shares of {sym}; available quantity is {available:g}')
            remaining = qty
            arr = lots[sym]
            while remaining > 1e-9 and arr:
                take = min(remaining, arr[0])
                arr[0] -= take
                remaining -= take
                if arr[0] <= 1e-9:
                    arr.pop(0)
        else:
            lots.setdefault(sym, []).append(qty)
        last_date[sym] = d


def _cashflows(sc, valuation_date):
    """Build the complete dated cash-flow stream used for Smallcase XIRR.

    XIRR starts from the earliest actual BUY transaction date automatically
    because calculate_xirr() uses the first dated cash flow as its base date.
    Every BUY is a negative investment cash flow and every SELL is a positive
    realized cash flow. Open holdings are handled by _summary(), which adds
    the current market value as the terminal cash flow on the valuation date.
    Therefore realized gains/losses from exits and unrealized gains/losses in
    currently held stocks are both reflected in the XIRR.
    """
    vd = _date(valuation_date, 'valuation date')
    flows = []
    for tx in _transactions(sc):
        d = _date(tx.get('transaction_date'), 'transaction date')
        if d > vd:
            continue
        gross = float(tx.get('gross_amount', 0) or 0)
        charges = float(tx.get('charges', 0) or 0)
        if tx.get('transaction_type') == 'BUY':
            # Investment outflow: purchase amount + buy charges.
            flows.append((d, -(gross + charges)))
        elif tx.get('transaction_type') == 'SELL':
            # Realized inflow: sale proceeds - sell charges.
            flows.append((d, gross - charges))
    return flows


def calculate_xirr(cashflows):
    """Return decimal XIRR or None using actual calendar days."""
    if not cashflows:
        return None
    flows = [(d if isinstance(d, dt.date) else _date(d, 'cashflow date'), float(a)) for d, a in cashflows if float(a) != 0]
    if not flows or not any(a < 0 for _, a in flows) or not any(a > 0 for _, a in flows):
        return None
    flows.sort(key=lambda x: x[0])
    base = flows[0][0]

    def npv(rate):
        if rate <= -1:
            return float('inf')
        return sum(a / ((1.0 + rate) ** ((d - base).days / 365.0)) for d, a in flows)

    # Scan a wide domain for a sign-changing bracket. This handles negative and very high XIRRs.
    points = [-0.999999, -0.9999, -0.999, -0.99, -0.95, -0.9, -0.75, -0.5, -0.25, -0.1, 0.0,
              0.1, 0.25, 0.5, 1, 2, 5, 10, 25, 50, 100, 250, 1000, 10000, 100000]
    prev_x, prev_y = points[0], npv(points[0])
    if abs(prev_y) < 1e-10:
        return prev_x
    bracket = None
    for x in points[1:]:
        y = npv(x)
        if abs(y) < 1e-10:
            return x
        if (prev_y < 0 < y) or (prev_y > 0 > y):
            bracket = (prev_x, x)
            break
        prev_x, prev_y = x, y
    if not bracket:
        # Newton with a few starting points can find roots that are outside the coarse bracket grid.
        for guess in (0.1, 0.5, 1.0, 2.0, 5.0, 10.0, -0.5, -0.9):
            r = guess
            for _ in range(100):
                if r <= -0.999999999:
                    break
                f = npv(r)
                denom = 1.0 + r
                df = sum(-((d - base).days / 365.0) * a / (denom ** (((d - base).days / 365.0) + 1.0)) for d, a in flows)
                if abs(df) < 1e-14:
                    break
                nr = r - f / df
                if nr <= -1 or not math.isfinite(nr) or abs(nr) > 1e9:
                    break
                if abs(nr - r) < 1e-10:
                    return nr
                r = nr
            if r > -1 and math.isfinite(r) and abs(npv(r)) < 1e-5:
                return r
        return None
    lo, hi = bracket
    flo = npv(lo)
    for _ in range(200):
        mid = (lo + hi) / 2.0
        fm = npv(mid)
        if abs(fm) < 1e-8 or abs(hi - lo) < 1e-10:
            return mid
        if (flo < 0 < fm) or (flo > 0 > fm):
            hi = mid
        else:
            lo = mid
            flo = fm
    return (lo + hi) / 2.0


def _market_price(symbol):
    if _FETCHER is None:
        return None, None, None, None
    try:
        q = _FETCHER.fetch_stock(symbol) or {}
        if not q.get('error') and q.get('price') is not None:
            price = float(q['price'])
            raw_change = q.get('change')
            raw_change_pct = q.get('change_pct')
            # Prefer the fetcher's explicit daily change. If unavailable,
            # derive it from previous close when supplied by the fetcher.
            if raw_change is None and q.get('previous_close') is not None:
                raw_change = price - float(q['previous_close'])
            if raw_change is None and raw_change_pct is not None and float(raw_change_pct) != -100:
                raw_change = price * float(raw_change_pct) / (100.0 + float(raw_change_pct))
            if raw_change_pct is None and q.get('previous_close') not in (None, 0):
                raw_change_pct = (float(raw_change or 0) / float(q['previous_close'])) * 100.0
            return (price, q.get('name') or symbol,
                    float(raw_change) if raw_change is not None else None,
                    float(raw_change_pct) if raw_change_pct is not None else None)
    except Exception:
        pass
    return None, None, None, None


def _holding_rows(sc):
    state = _position_state(sc)
    rows = []
    for sym, st in state.items():
        qty = st['remaining_qty']
        price, live_name, day_change, day_change_pct = _market_price(sym) if qty > 0 else (None, None, None, None)
        current = price * qty if price is not None else None
        today_change = day_change * qty if day_change is not None else None
        prev_value = ((price - day_change) * qty) if price is not None and day_change is not None else None
        unrealized = current - st['invested_remaining'] if current is not None else None
        total_sold = st['sell_proceeds']
        total_pnl = st['realized_pnl'] + (unrealized or 0.0)
        latest_tx = max((tx for tx in _transactions(sc) if _symbol(tx.get('symbol')) == sym), key=lambda tx: (tx.get('created_at', ''), tx.get('transaction_date', '')), default=None)
        last_activity = latest_tx.get('created_at') if latest_tx else None
        rows.append({
            'symbol': sym,
            'stock_name': live_name or st['name'] or sym,
            'quantity': round(qty, 6),
            'average_buy_price': round(st['avg_buy'], 2),
            'invested': round(st['invested_remaining'], 2),
            'current_price': round(price, 2) if price is not None else None,
            'current_value': round(current, 2) if current is not None else None,
            'today_change': round(today_change, 2) if today_change is not None else None,
            'today_change_pct': round((today_change / prev_value) * 100, 2) if today_change is not None and prev_value not in (None, 0) else (round(day_change_pct, 2) if day_change_pct is not None else None),
            'previous_close_value': round(prev_value, 2) if prev_value is not None else None,
            'realized_pnl': round(st['realized_pnl'], 2),
            'unrealized_pnl': round(unrealized, 2) if unrealized is not None else None,
            'total_pnl': round(total_pnl, 2),
            # Return remains available after a full exit. Use the total historical
            # buy cost as the denominator instead of the remaining invested amount.
            'return_pct': round(total_pnl / st['buy_cost'] * 100, 2) if st['buy_cost'] else 0,
            'status': 'ACTIVE' if qty > 1e-9 else 'CLOSED',
            'last_activity': last_activity,
            'total_bought': round(st['buys'], 6),
            'total_sold': round(st['sells'], 6),
            'market_price_available': price is not None,
        })
    return rows


def _summary(sc, valuation_date=None):
    valuation_date = valuation_date or _today()
    rows = _holding_rows(sc)
    txs = _transactions(sc)
    buy_total = sum(float(t.get('gross_amount', 0) or 0) + float(t.get('charges', 0) or 0) for t in txs if t.get('transaction_type') == 'BUY')
    sell_total = sum(float(t.get('gross_amount', 0) or 0) - float(t.get('charges', 0) or 0) for t in txs if t.get('transaction_type') == 'SELL')
    current_value = sum(r['current_value'] or 0 for r in rows if r['status'] == 'ACTIVE')
    realized = sum(r['realized_pnl'] for r in rows)
    unrealized = sum(r['unrealized_pnl'] or 0 for r in rows)
    total_pnl = sell_total + current_value - buy_total
    day_changes = [r['today_change'] for r in rows if r['status'] == 'ACTIVE' and r['today_change'] is not None]
    previous_values = [r['previous_close_value'] for r in rows if r['status'] == 'ACTIVE' and r['previous_close_value'] is not None]
    today_change = sum(day_changes) if day_changes else None
    previous_day_value = sum(previous_values) if previous_values else None
    today_change_pct = (today_change / previous_day_value * 100) if today_change is not None and previous_day_value else None
    # XIRR uses every historical BUY/SELL cash flow plus the current market
    # value of still-open holdings at the valuation date. This captures both
    # realized P&L (SELL proceeds) and unrealized P&L (terminal market value).
    xirr_flows = _cashflows(sc, valuation_date)
    if active_positions := [r for r in rows if r['status'] == 'ACTIVE']:
        if current_value > 0:
            xirr_flows.append((_date(valuation_date, 'valuation date'), current_value))
    xirr = calculate_xirr(xirr_flows)
    active = sum(r['status'] == 'ACTIVE' for r in rows)
    total_symbols = len(rows)
    has_sell = any(t.get('transaction_type') == 'SELL' for t in txs)
    if total_symbols == 0:
        status = 'CLOSED'
    elif active == 0:
        status = 'CLOSED'
    else:
        # A current active position is ACTIVE; PARTIALLY EXITED means at least one
        # current holding has a quantity below the historical quantity bought.
        partial = any(r['status'] == 'ACTIVE' and r['total_sold'] > 0 for r in rows)
        status = 'PARTIALLY EXITED' if partial else 'ACTIVE'
    # Returns Since is anchored to the first actual BUY transaction date.
    # This must not depend on the Smallcase creation/investment-date field.
    buy_dates = [t.get('transaction_date') for t in txs
                 if t.get('transaction_type') == 'BUY' and t.get('transaction_date')]
    first_buy = min((t for t in txs if t.get('transaction_type') == 'BUY' and t.get('transaction_date')), key=lambda t: (t.get('transaction_date',''), t.get('created_at','')), default=None)
    first_stock_date = first_buy.get('transaction_date') if first_buy else sc.get('investment_date')
    returns_since = (first_stock_date + ' ' + str(first_buy.get('created_at',''))[11:19]) if first_buy and first_buy.get('created_at') else first_stock_date

    return {
        'id': sc['id'], 'name': sc['name'], 'description': sc.get('description', ''), 'category': sc.get('category', ''),
        'benchmark': sc.get('benchmark', ''), 'notes': sc.get('notes', ''), 'investment_date': first_stock_date, 'returns_since': returns_since,
        'created_at': sc.get('created_at'), 'updated_at': sc.get('updated_at'), 'closed_at': sc.get('closed_at'),
        'status': status, 'stocks_count': total_symbols, 'active_positions': active, 'closed_positions': total_symbols - active,
        'total_invested': round(buy_total, 2), 'current_value': round(current_value, 2), 'realized_pnl': round(realized, 2),
        'unrealized_pnl': round(unrealized, 2), 'total_pnl': round(total_pnl, 2),
        'return_pct': round(total_pnl / buy_total * 100, 2) if buy_total else 0,
        'xirr': round(xirr * 100, 2) if xirr is not None else None, 'xirr_as_of': valuation_date,
        'today_change': round(today_change, 2) if today_change is not None else None,
        'today_change_pct': round(today_change_pct, 2) if today_change_pct is not None else None,
        'market_price_available': all(r['market_price_available'] for r in rows if r['status'] == 'ACTIVE'),
    }


def _all_cashflows(data, valuation_date):
    flows = []
    vd = _date(valuation_date, 'valuation date')
    for sc in data.values():
        flows.extend(_cashflows(sc, valuation_date))
        current = sum((r['current_value'] or 0) for r in _holding_rows(sc) if r['status'] == 'ACTIVE')
        if current > 0:
            flows.append((vd, current))
    return flows


def _guard(f):
    """Apply the existing app login_required decorator at request time.

    smallcases.py is imported before init_smallcases() is called, so the
    authentication decorator must not be evaluated during module import.
    """
    from functools import wraps

    @wraps(f)
    def wrapper(*args, **kwargs):
        if _LOGIN_REQUIRED is None:
            from flask import jsonify
            return jsonify({
                "error": "Smallcases module is not initialized"
            }), 500

        protected = _LOGIN_REQUIRED(f)
        return protected(*args, **kwargs)

    return wrapper


@smallcases_bp.route('/smallcases', methods=['GET'])
@_guard
def page():
    return render_template('index.html', initial_tab='smallcases')


@smallcases_bp.route('/api/smallcases', methods=['GET'])
@_guard
def list_smallcases():
    data = _load()
    valuation = request.args.get('valuation_date') or _today()
    rows = [_summary(sc, valuation) for sc in data.values()]
    rows.sort(key=lambda x: x.get('created_at') or '', reverse=True)
    total_invested = sum(x['total_invested'] for x in rows)
    current_value = sum(x['current_value'] for x in rows)
    total_pnl = sum(x['total_pnl'] for x in rows)
    day_values = [x for x in rows if x.get('active_positions', 0) > 0 and x.get('today_change') is not None]
    overall_today_change = sum(x['today_change'] for x in day_values) if day_values else None
    overall_previous_value = sum(
        x['current_value'] - x['today_change'] for x in day_values
        if x.get('current_value') is not None and x.get('today_change') is not None
    ) if day_values else None
    overall_today_change_pct = (overall_today_change / overall_previous_value * 100) if overall_today_change is not None and overall_previous_value else None
    overall_xirr = calculate_xirr(_all_cashflows(data, valuation))
    return jsonify({'smallcases': rows, 'summary': {
        'total_smallcases': len(rows),
        'active_smallcases': sum(x['active_positions'] > 0 for x in rows),
        'partially_exited_smallcases': sum(x['status'] == 'PARTIALLY EXITED' for x in rows),
        'closed_smallcases': sum(x['status'] == 'CLOSED' for x in rows),
        'total_invested': round(total_invested, 2), 'current_value': round(current_value, 2),
        'total_pnl': round(total_pnl, 2), 'return_pct': round(total_pnl / total_invested * 100, 2) if total_invested else 0,
        'today_change': round(overall_today_change, 2) if overall_today_change is not None else None,
        'today_change_pct': round(overall_today_change_pct, 2) if overall_today_change_pct is not None else None,
        'overall_xirr': round(overall_xirr * 100, 2) if overall_xirr is not None else None,
        'xirr_as_of': datetime.now().isoformat(timespec='minutes')
    }})


@smallcases_bp.route('/api/smallcases', methods=['POST'])
@_guard
def create_smallcase():
    try:
        b = request.get_json(force=True) or {}
        name = str(b.get('name', '')).strip()
        if not name:
            raise ValueError('Smallcase name is required')
        data = _load()
        if any(str(sc.get('name', '')).strip().lower() == name.lower() for sc in data.values()):
            return jsonify({'error': 'Smallcase already exists'}), 400
        investment_date = str(b.get('investment_date') or _today())
        _date(investment_date, 'investment date')
        now = _now()
        sc = {'id': _id('sc'), 'name': name, 'description': str(b.get('description', '')).strip(),
              'category': str(b.get('category', '')).strip(), 'benchmark': str(b.get('benchmark', '')).strip(),
              'notes': str(b.get('notes', '')).strip(), 'investment_date': investment_date,
              'created_at': now, 'updated_at': now, 'closed_at': None, 'transactions': [], 'audit_logs': []}
        _audit(sc, 'SMALLCASE_CREATED', 'smallcase', sc['id'], new={k: sc[k] for k in ('name','description','category','benchmark','investment_date')})
        data[sc['id']] = sc
        _save(data)
        return jsonify({'success': True, 'smallcase': _summary(sc)}), 201
    except ValueError as e:
        return jsonify({'error': str(e)}), 400
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@smallcases_bp.route('/api/smallcases/<identifier>', methods=['GET'])
@_guard
def detail(identifier):
    data = _load(); key, sc = _case_key(data, identifier)
    if not sc: return jsonify({'error': 'Smallcase not found'}), 404
    valuation = request.args.get('valuation_date') or _today()
    out = _summary(sc, valuation)
    out.update({'holdings': _holding_rows(sc), 'transactions': sorted(_transactions(sc), key=lambda x: (x.get('transaction_date',''), x.get('created_at','')), reverse=True), 'audit_logs': sc.get('audit_logs', [])})
    return jsonify(out)


@smallcases_bp.route('/api/smallcases/<identifier>', methods=['PUT'])
@_guard
def update_smallcase(identifier):
    data = _load(); key, sc = _case_key(data, identifier)
    if not sc: return jsonify({'error': 'Smallcase not found'}), 404
    b = request.get_json(force=True) or {}
    new_name = str(b.get('name', sc['name'])).strip()
    if not new_name: return jsonify({'error': 'Smallcase name is required'}), 400
    for other_key, other in data.items():
        if other_key != key and str(other.get('name','')).strip().lower() == new_name.lower():
            return jsonify({'error': 'Smallcase name already exists'}), 400
    fields = ('name','description','category','benchmark','notes','investment_date')
    old = {f: sc.get(f, '') for f in fields}
    for f in fields:
        if f in b:
            if f == 'investment_date': _date(b[f], f)
            sc[f] = str(b[f]).strip()
    sc['updated_at'] = _now()
    _audit(sc, 'SMALLCASE_UPDATED', 'smallcase', sc['id'], old=old, new={f: sc.get(f,'') for f in fields})
    if new_name != key:
        del data[key]
        data[sc['id']] = sc
    _save(data)
    return jsonify({'success': True, 'smallcase': _summary(sc)})


@smallcases_bp.route('/api/smallcases/<identifier>', methods=['DELETE'])
@_guard
def delete_smallcase(identifier):
    data = _load(); key, sc = _case_key(data, identifier)
    if not sc: return jsonify({'error': 'Smallcase not found'}), 404
    # Delete is deliberately blocked once financial history exists. This preserves the audit/transaction record.
    if _transactions(sc):
        return jsonify({'error': 'Financial history cannot be deleted. Fully exit the Smallcase and keep it as CLOSED.'}), 400
    _audit(sc, 'SMALLCASE_DELETED', 'smallcase', sc['id'], old={'name': sc['name']}, details='Empty Smallcase deleted')
    del data[key]; _save(data)
    return jsonify({'success': True})


@smallcases_bp.route('/api/smallcases/<identifier>/transactions', methods=['POST'])
@_guard
def add_transaction(identifier):
    try:
        data = _load(); key, sc = _case_key(data, identifier)
        if not sc: return jsonify({'error': 'Smallcase not found'}), 404
        b = request.get_json(force=True) or {}
        typ = str(b.get('transaction_type', 'BUY')).upper()
        sym = _symbol(b.get('symbol'))
        qty = _num(b.get('quantity'), 'quantity')
        price = _num(b.get('price'), 'price')
        d = _date(b.get('transaction_date') or b.get('buy_date'), 'transaction date')
        charges = _num(b.get('charges', 0), 'charges', allow_zero=True)
        if typ not in ('BUY','SELL'): raise ValueError('transaction_type must be BUY or SELL')
        # Always build the candidate transaction list before validation.
        # The previous implementation initialized `candidate` only for SELL
        # transactions, which caused BUY requests to fail with:
        # UnboundLocalError: cannot access local variable 'candidate'...
        candidate = list(_transactions(sc))
        gross = qty * price
        if typ == 'SELL' and gross - charges < -1e-9:
            raise ValueError('Charges cannot exceed sale proceeds')
        stock_name = str(b.get('stock_name') or sym)
        if _FETCHER is not None:
            try:
                q = _FETCHER.fetch_stock(sym) or {}
                if not q.get('error'):
                    stock_name = q.get('name') or stock_name
            except Exception:
                pass
        tx = {'id': _id('tx'), 'smallcase_id': sc['id'], 'stock_id': sym, 'stock_symbol': sym, 'symbol': sym,
              'stock_name': stock_name, 'transaction_type': typ, 'quantity': round(qty, 6), 'price': round(price, 6),
              'transaction_date': d.strftime('%Y-%m-%d'), 'gross_amount': round(gross, 2), 'charges': round(charges, 2),
              'net_amount': round(gross + charges if typ == 'BUY' else gross - charges, 2), 'notes': str(b.get('notes','')).strip(),
              'created_at': _now(), 'updated_at': _now(), 'created_by': session.get('username')}
        candidate.append(tx)
        _validate_transactions(sc, candidate)
        sc['transactions'] = candidate
        sc['updated_at'] = _now()
        _audit(sc, typ + '_TRANSACTION_ADDED', 'transaction', tx['id'], new=tx)
        _save(data)
        return jsonify({'success': True, 'transaction': tx, 'smallcase': _summary(sc)}), 201
    except ValueError as e:
        return jsonify({'error': str(e)}), 400
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@smallcases_bp.route('/api/smallcases/<identifier>/import', methods=['POST'])
@_guard
def import_transactions(identifier):
    """Import BUY/SELL transactions from an Excel workbook.

    Imports are additive and can be repeated. A row is skipped when an exact
    transaction already exists or when the row fails validation. Historical
    transactions are never overwritten.
    """
    try:
        data = _load(); key, sc = _case_key(data, identifier)
        if not sc:
            return jsonify({'error': 'Smallcase not found'}), 404
        upload = request.files.get('file')
        if upload is None or not upload.filename:
            return jsonify({'error': 'Please select an Excel file'}), 400
        filename = upload.filename.lower()
        if not filename.endswith(('.xlsx', '.xlsm')):
            return jsonify({'error': 'Only .xlsx or .xlsm Excel files are supported'}), 400

        try:
            from openpyxl import load_workbook
        except Exception:
            return jsonify({'error': 'Excel import requires openpyxl on the server'}), 500

        wb = load_workbook(upload, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            return jsonify({'error': 'Excel file is empty'}), 400

        headers = {str(v or '').strip().lower(): i for i, v in enumerate(header)}
        required = ['symbol', 'date', 'buy/sell', 'quantity', 'buy price', 'sell price']
        missing = [h for h in required if h not in headers]
        if missing:
            return jsonify({'error': 'Missing required columns: ' + ', '.join(missing)}), 400

        existing = list(_transactions(sc))
        fingerprints = {_import_fingerprint(tx) for tx in existing}
        candidate = list(existing)
        imported = []
        already_imported = []
        errored = []

        for row_no, values in enumerate(rows, start=2):
            if not any(v is not None and str(v).strip() for v in values):
                continue
            try:
                raw_symbol = values[headers['symbol']] if headers['symbol'] < len(values) else None
                sym = _symbol(raw_symbol)
                tx_date = _import_date(values[headers['date']] if headers['date'] < len(values) else None)
                typ = str(values[headers['buy/sell']] if headers['buy/sell'] < len(values) else '').strip().upper()
                if typ not in ('BUY', 'SELL'):
                    raise ValueError('Buy/Sell must be BUY or SELL')
                qty = _num(values[headers['quantity']] if headers['quantity'] < len(values) else None, 'quantity')
                buy_price = values[headers['buy price']] if headers['buy price'] < len(values) else None
                sell_price = values[headers['sell price']] if headers['sell price'] < len(values) else None
                raw_price = buy_price if typ == 'BUY' else sell_price
                if raw_price is None or str(raw_price).strip() == '':
                    raise ValueError(f'{typ} Price is required')
                price = _num(raw_price, f'{typ} Price')

                tx = {
                    'id': _id('tx'), 'smallcase_id': sc['id'], 'stock_id': sym,
                    'stock_symbol': sym, 'symbol': sym, 'stock_name': sym,
                    'transaction_type': typ, 'quantity': round(qty, 6),
                    'price': round(price, 6), 'transaction_date': tx_date.strftime('%Y-%m-%d'),
                    'gross_amount': round(qty * price, 2), 'charges': 0.0,
                    'net_amount': round(qty * price, 2), 'notes': 'Imported from Excel',
                    'created_at': _now(), 'updated_at': _now(),
                    'created_by': session.get('username')
                }
                fp = _import_fingerprint(tx)
                if fp in fingerprints:
                    already_imported.append({
                        'row': row_no,
                        'symbol': sym,
                        'type': typ,
                        'date': tx['transaction_date'],
                        'quantity': tx['quantity'],
                        'price': tx['price'],
                        'reason': 'Already imported/existing transaction'
                    })
                    continue
                trial = candidate + [tx]
                _validate_transactions(sc, trial)
                candidate = trial
                fingerprints.add(fp)
                imported.append({'row': row_no, 'symbol': sym, 'type': typ, 'date': tx['transaction_date'], 'quantity': tx['quantity'], 'price': tx['price']})
            except Exception as e:
                errored.append({
                    'row': row_no,
                    'symbol': str(values[headers['symbol']]).strip() if headers['symbol'] < len(values) and values[headers['symbol']] is not None else '',
                    'reason': str(e)
                })

        if imported:
            sc['transactions'] = candidate
            sc['updated_at'] = _now()
            for item, tx in zip(imported, [t for t in candidate if t not in existing]):
                _audit(sc, tx['transaction_type'] + '_TRANSACTION_IMPORTED', 'transaction', tx['id'], new=tx, details=f"Imported Excel row {item['row']}")
            _save(data)

        return jsonify({
            'success': True, 'filename': upload.filename,
            'total_rows': len(imported) + len(already_imported) + len(errored),
            'imported_count': len(imported),
            'already_imported_count': len(already_imported),
            'errored_count': len(errored),
            'not_imported_count': len(already_imported) + len(errored),
            'imported': imported,
            'already_imported': already_imported,
            'errored': errored,
            # Backward-compatible alias for clients expecting not_imported.
            'not_imported': errored,
            'smallcase': _summary(sc),
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@smallcases_bp.route('/api/smallcases/<identifier>/history', methods=['GET'])
@_guard
def history(identifier):
    data = _load(); _, sc = _case_key(data, identifier)
    if not sc: return jsonify({'error': 'Smallcase not found'}), 404
    return jsonify(sc.get('audit_logs', []))


@smallcases_bp.route('/api/smallcases/<identifier>/xirr', methods=['GET'])
@_guard
def xirr(identifier):
    data = _load(); _, sc = _case_key(data, identifier)
    if not sc: return jsonify({'error': 'Smallcase not found'}), 404
    valuation = request.args.get('valuation_date') or _today()
    flows = _cashflows(sc, valuation)
    current = sum((r['current_value'] or 0) for r in _holding_rows(sc) if r['status'] == 'ACTIVE')
    if current > 0: flows.append((_date(valuation, 'valuation date'), current))
    rate = calculate_xirr(flows)
    return jsonify({'smallcase_id': sc['id'], 'smallcase_name': sc['name'], 'valuation_date': valuation,
                    'xirr': round(rate * 100, 2) if rate is not None else None,
                    'xirr_decimal': rate, 'cashflows': [{'date': d.strftime('%Y-%m-%d'), 'amount': round(a,2)} for d,a in sorted(flows)]})


@smallcases_bp.route('/api/smallcases/overall-xirr', methods=['GET'])
@_guard
def overall_xirr():
    data = _load(); valuation = request.args.get('valuation_date') or _today()
    flows = _all_cashflows(data, valuation)
    rate = calculate_xirr(flows)
    return jsonify({'valuation_date': valuation, 'xirr': round(rate*100,2) if rate is not None else None,
                    'xirr_decimal': rate, 'cashflows': [{'date': d.strftime('%Y-%m-%d'), 'amount': round(a,2)} for d,a in sorted(flows)]})


@smallcases_bp.route('/api/smallcases/<identifier>/close', methods=['POST'])
@_guard
def close_smallcase(identifier):
    data = _load(); _, sc = _case_key(data, identifier)
    if not sc: return jsonify({'error': 'Smallcase not found'}), 404
    if any(r['status'] == 'ACTIVE' for r in _holding_rows(sc)):
        return jsonify({'error': 'Exit all active holdings before closing the Smallcase'}), 400
    sc['closed_at'] = _now(); sc['updated_at'] = _now()
    _audit(sc, 'SMALLCASE_CLOSED', 'smallcase', sc['id'], new={'status':'CLOSED'})
    _save(data)
    return jsonify({'success': True, 'smallcase': _summary(sc)})


@smallcases_bp.route('/api/smallcases/<identifier>/performance', methods=['GET'])
@_guard
def performance(identifier):
    data = _load(); _, sc = _case_key(data, identifier)
    if not sc: return jsonify({'error': 'Smallcase not found'}), 404
    # Only show actual transaction dates. No artificial historical prices are generated.
    points = []
    for d in sorted(set(t.get('transaction_date') for t in _transactions(sc) if t.get('transaction_date'))):
        invested = sum(float(t.get('gross_amount',0) or 0) + float(t.get('charges',0) or 0) for t in _transactions(sc) if t.get('transaction_type') == 'BUY' and t.get('transaction_date') <= d)
        proceeds = sum(float(t.get('gross_amount',0) or 0) - float(t.get('charges',0) or 0) for t in _transactions(sc) if t.get('transaction_type') == 'SELL' and t.get('transaction_date') <= d)
        points.append({'date': d, 'invested_capital': round(invested,2), 'realized_proceeds': round(proceeds,2), 'portfolio_value': None})
    return jsonify({'points': points, 'note': 'Historical portfolio value is not fabricated when historical market prices are unavailable.'})
