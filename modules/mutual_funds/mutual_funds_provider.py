"""AMFI NAV feed access and parsing for the Mutual Funds module."""

from __future__ import annotations

import datetime as dt
from io import BytesIO
import re
import threading
import time
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook


AMFI_NAV_PAGE = "https://www.amfiindia.com/net-asset-value/nav-download"
LEGACY_NAV_URLS = (
    "https://portal.amfiindia.com/spages/NAVAll.txt",
    "https://www.amfiindia.com/spages/NAVAll.txt",
)
AMFI_NAV_SOURCE = "https://www.amfiindia.com/net-asset-value/nav-download"
MFAPI_NAV_SOURCE = "https://api.mfapi.in/"
TIMEOUT = (6, 25)
HEADERS = {
    "User-Agent": "CapitalDesk Mutual Funds/1.0",
    "Accept": "text/plain,text/html,application/octet-stream,*/*",
}


class AMFIError(RuntimeError):
    """Raised when AMFI cannot provide a parseable response."""


def _decode(response):
    content_type = str(response.headers.get("Content-Type", "")).lower()
    if "charset=" not in content_type:
        try:
            return response.content.decode("utf-8-sig")
        except UnicodeDecodeError:
            pass
    response.encoding = response.encoding or "utf-8"
    return response.text.lstrip("\ufeff")


def _discover_report_url(label="Complete NAV Report", old=False):
    """Discover AMFI's download URL from its current NAV page markup."""
    try:
        response = requests.get(AMFI_NAV_PAGE, headers=HEADERS, timeout=TIMEOUT)
        response.raise_for_status()
    except requests.RequestException as exc:
        raise AMFIError("AMFI's NAV download page could not be reached.") from exc

    soup = BeautifulSoup(_decode(response), "html.parser")
    candidates = []
    for element in soup.find_all(["a", "button", "input"]):
        text = " ".join([
            element.get_text(" ", strip=True),
            str(element.get("value") or ""),
            str(element.get("title") or ""),
            str(element.get("aria-label") or ""),
        ])
        onclick = str(element.get("onclick") or "")
        href = str(element.get("href") or "")
        attributes = " ".join(
            str(attr_value)
            for key, raw_value in element.attrs.items() if key.startswith("data-")
            for attr_value in (raw_value if isinstance(raw_value, list) else [raw_value])
        )
        vicinity = " ".join([
            text,
            str(element.parent.get_text(" ", strip=True) if element.parent else ""),
        ]).lower()
        if label.lower() not in text.lower() and label.lower() not in onclick.lower():
            continue
        marks_old = "old version" in vicinity or "old format" in vicinity
        if marks_old != old:
            continue
        matches = re.findall(r"(?:https?://[^\s'\"]+|/[A-Za-z0-9_./?&=%+-]+)", " ".join((href, onclick, attributes)))
        match = max(matches, key=len) if matches else None
        if match:
            candidates.append(urljoin(response.url, match.rstrip(");}")))
    if not candidates:
        raise AMFIError("AMFI did not expose a downloadable NAV report link.")
    return candidates[0]


def _parse_date(value):
    raw = str(value or "").strip()
    for fmt in ("%d-%b-%Y", "%d-%B-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    return None


def _category_header(value):
    raw = value.strip()
    low = raw.lower()
    if "scheme" not in low and "interval fund" not in low:
        return None
    if "open ended" in low:
        scheme_type = "Open Ended"
    elif "close ended" in low or "closed ended" in low:
        scheme_type = "Close Ended"
    elif "interval fund" in low:
        scheme_type = "Interval Fund"
    else:
        return None
    category = ""
    match = re.search(r"\((.*)\)", raw)
    if match:
        category = re.sub(r"^\s*(?:equity|debt|hybrid|solution oriented|other)\s+scheme\s*[-–:]?\s*", "", match.group(1), flags=re.I).strip()
    return scheme_type, category or scheme_type


def _split_data_line(raw):
    delimiter = ";" if raw.count(";") >= raw.count("|") else "|"
    fields = [part.strip() for part in raw.split(delimiter)]
    while fields and not fields[-1]:
        fields.pop()
    return fields


def parse_nav_report(text):
    """Normalize AMFI's old and current semicolon-delimited NAV reports."""
    schemes = []
    amc = "Unknown AMC"
    category = "Unclassified"
    scheme_type = "Unknown"
    for line in (text or "").splitlines():
        raw = line.strip().strip("\ufeff")
        if not raw:
            continue
        if ";" not in raw and "|" not in raw:
            header = _category_header(raw)
            if header:
                scheme_type, category = header
            elif "mutual fund" in raw.lower() or "asset management" in raw.lower():
                amc = raw
            continue

        fields = _split_data_line(raw)
        if not fields or not re.fullmatch(r"\d{3,8}", fields[0]):
            continue
        if len(fields) < 6:
            continue
        nav_date = _parse_date(fields[-1])
        try:
            nav = float(fields[-2].replace(",", ""))
        except (TypeError, ValueError):
            continue
        if not nav_date or nav < 0 or not fields[3]:
            continue

        plan = fields[4] if len(fields) >= 8 else ""
        option = fields[5] if len(fields) >= 8 else ""
        name = fields[3]
        low_name = name.lower()
        if not plan:
            plan = "Direct" if re.search(r"\bdirect\b", low_name) else ("Regular" if re.search(r"\bregular\b", low_name) else "Not stated")
        else:
            plan = "Direct" if "direct" in plan.lower() else ("Regular" if "regular" in plan.lower() else plan)
        if not option:
            if "idcw" in low_name or "dividend" in low_name:
                option = "IDCW"
            elif "growth" in low_name:
                option = "Growth"
            else:
                option = "Not stated"
        else:
            option = "IDCW" if "idcw" in option.lower() else ("Growth" if "growth" in option.lower() else option)

        schemes.append({
            "scheme_code": fields[0],
            "isin": fields[1] if len(fields) > 1 and fields[1] != "-" else "",
            "isin_reinvestment": fields[2] if len(fields) > 2 and fields[2] != "-" else "",
            "name": name,
            "amc": amc,
            "category": category,
            "scheme_type": scheme_type,
            "plan": plan,
            "option": option,
            "nav": nav,
            "nav_date": nav_date.isoformat(),
        })
    if not schemes:
        raise AMFIError("AMFI's NAV report was downloaded, but its format was not recognized.")
    schemes.sort(key=lambda row: (row["amc"].casefold(), row["name"].casefold(), row["scheme_code"]))
    return schemes


_LATEST_NAVS_CACHE = None
_LATEST_NAVS_LOCK = threading.Lock()
_LATEST_NAVS_TTL = 300  # Cache for 5 minutes

_SCHEME_NAV_CACHE = {}
_SCHEME_NAV_LOCK = threading.Lock()
_SCHEME_NAV_TTL = 21600  # Cache scheme history for 6 hours


def fetch_latest_navs(force_refresh=False):
    """Fetch the complete NAV report; cached for 5 minutes to avoid redundant downloads."""
    global _LATEST_NAVS_CACHE
    now = time.monotonic()
    if not force_refresh:
        with _LATEST_NAVS_LOCK:
            if _LATEST_NAVS_CACHE and _LATEST_NAVS_CACHE[0] > now:
                return _LATEST_NAVS_CACHE[1]

    errors = []
    urls = list(LEGACY_NAV_URLS)
    try:
        discovered = _discover_report_url()
        if discovered not in urls:
            urls.append(discovered)
    except AMFIError as exc:
        errors.append(str(exc))

    for url in urls:
        try:
            response = requests.get(url, headers=HEADERS, timeout=(3, 6))
            response.raise_for_status()
            content = _decode(response)
            schemes = parse_nav_report(content)
            nav_as_of = max((row["nav_date"] for row in schemes), default=None)
            if nav_as_of and (dt.datetime.now().date() - dt.date.fromisoformat(nav_as_of)).days > 30:
                raise AMFIError(f"AMFI's latest published NAV date is {nav_as_of}; the report appears out of date.")
            feed = {
                "schemes": schemes,
                "source": url,
                "source_label": "AMFI latest NAV report",
                "fetched_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
                "nav_as_of": nav_as_of,
            }
            with _LATEST_NAVS_LOCK:
                _LATEST_NAVS_CACHE = (now + _LATEST_NAVS_TTL, feed)
            return feed
        except (requests.RequestException, AMFIError) as exc:
            errors.append(str(exc))

    with _LATEST_NAVS_LOCK:
        if _LATEST_NAVS_CACHE:
            return _LATEST_NAVS_CACHE[1]

    raise AMFIError("AMFI's latest NAV report is unavailable or unreadable. " + " ".join(errors[-2:]))


def _parse_history_rows(text):
    result = {}
    field_rows = []
    for line in (text or "").splitlines():
        raw = line.strip().strip("\ufeff")
        if not raw or (";" not in raw and "|" not in raw):
            continue
        field_rows.append(_split_data_line(raw))
    def add_row(fields):
        if len(fields) < 4 or not re.fullmatch(r"\d{3,8}", fields[0]):
            return
        nav_date = _parse_date(fields[-1])
        # The AMFI history endpoint uses this legacy row order:
        # Code;Name;ISIN;ISIN;NAV;Repurchase;Sale;Date.
        nav_field = fields[-4] if len(fields) >= 8 and not re.fullmatch(r"INF[A-Z0-9]{8,}", fields[1], flags=re.I) else fields[-2]
        try:
            nav = float(nav_field.replace(",", ""))
        except (TypeError, ValueError):
            return
        if nav_date and nav > 0:
            result.setdefault(fields[0], []).append({"date": nav_date.isoformat(), "nav": nav})
    for fields in field_rows:
        add_row(fields)
    if not result and "<tr" in (text or "").lower():
        soup = BeautifulSoup(text, "html.parser")
        for row in soup.find_all("tr"):
            add_row([cell.get_text(" ", strip=True) for cell in row.find_all(["td", "th"])])
    return result


def fetch_scheme_history(scheme_code, start_date, end_date, timeout=None):
    """Fetch scheme-specific historical NAVs from MFapi.in as an AMFI fallback."""
    code = str(scheme_code).strip()
    if not re.fullmatch(r"\d{3,8}", code):
        raise AMFIError("A valid scheme code is required for historical NAV fallback.")
    try:
        response = requests.get(
            f"{MFAPI_NAV_SOURCE}mf/{code}",
            params={"startDate": start_date.isoformat(), "endDate": end_date.isoformat()},
            headers={"User-Agent": HEADERS["User-Agent"], "Accept": "application/json"},
            timeout=timeout or TIMEOUT,
        )
        response.raise_for_status()
        payload = response.json()
    except (requests.RequestException, ValueError) as exc:
        raise AMFIError("MFapi.in could not provide fallback NAV history.") from exc

    now = time.monotonic()
    cached_records = None
    with _SCHEME_NAV_LOCK:
        cached = _SCHEME_NAV_CACHE.get(code)
        if cached and cached[0] > now:
            cached_records = cached[1]

    if cached_records is None:
        try:
            response = requests.get(
                f"{MFAPI_NAV_SOURCE}mf/{code}",
                headers={"User-Agent": HEADERS["User-Agent"], "Accept": "application/json"},
                timeout=TIMEOUT,
            )
            response.raise_for_status()
            payload = response.json()
        except (requests.RequestException, ValueError) as exc:
            raise AMFIError("MFapi.in could not provide fallback NAV history.") from exc

        metadata = payload.get("meta") if isinstance(payload, dict) else None
        if not isinstance(payload, dict) or payload.get("status") != "SUCCESS" or not isinstance(metadata, dict) or str(metadata.get("scheme_code", "")) != code:
            raise AMFIError("MFapi.in returned a failed response or a different scheme.")
        records = payload.get("data")
        if not isinstance(records, list):
            raise AMFIError("MFapi.in returned an invalid NAV history response.")

        all_navs = []
        for record in records:
            if not isinstance(record, dict):
                continue
            nav_date = _parse_date(record.get("date"))
            try:
                nav = float(str(record.get("nav", "")).replace(",", ""))
            except (TypeError, ValueError):
                continue
            if nav_date and nav > 0:
                all_navs.append({"date": nav_date.isoformat(), "nav": nav, "date_obj": nav_date})
        all_navs.sort(key=lambda row: row["date"])
        cached_records = all_navs
        with _SCHEME_NAV_LOCK:
            _SCHEME_NAV_CACHE[code] = (now + _SCHEME_NAV_TTL, cached_records)

    navs = [
        {"date": row["date"], "nav": row["nav"]}
        for row in cached_records
        if start_date <= row["date_obj"] <= end_date
    ]
    if not navs:
        raise AMFIError("MFapi.in returned no NAV rows for this scheme and date range.")
    return {code: navs}



_OFFICIAL_FACTS_CACHE = {}
_OFFICIAL_FACTS_LOCK = threading.Lock()


def fetch_official_fund_facts(scheme):
    code = str(scheme.get("scheme_code") or "")
    now = time.monotonic()
    with _OFFICIAL_FACTS_LOCK:
        cached = _OFFICIAL_FACTS_CACHE.get(code)
        if cached and cached[0] > now:
            return cached[1]
    facts = _fetch_official_fund_facts_uncached(scheme)
    with _OFFICIAL_FACTS_LOCK:
        _OFFICIAL_FACTS_CACHE[code] = (now + 21600, facts)
    return facts


def _fetch_official_fund_facts_uncached(scheme):
    """Fetch scheme facts from configured fund sources; never infer missing values."""
    amc = str(scheme.get("amc") or "").lower()
    scheme_name = str(scheme.get("name") or "")
    name = scheme_name.lower()
    is_sbi = "sbi" in amc and "mutual" in amc
    is_motilal = "motilal oswal" in amc
    is_hdfc = "hdfc" in amc
    if is_sbi:
        return _fetch_angel_one_fund_facts(scheme)
    if not is_motilal and not is_hdfc:
        return {}
    if is_motilal:
        slug_name = re.sub(r"^motilal oswal\s*", "", scheme_name.split(" - ")[0], flags=re.I).strip().lower()
        url = "https://www.motilaloswalmf.com/mutual-funds/motilal-oswal-" + re.sub(r"[^a-z0-9]+", "-", slug_name).strip("-")
    else:
        base_name = re.sub(r"\s*-\s*(direct|regular).*?$", "", scheme_name, flags=re.I).strip()
        slug = re.sub(r"[^a-z0-9]+", "-", base_name.lower()).strip("-")
        plan = "regular" if "regular" in str(scheme.get("plan") or "").lower() else "direct"
        url = f"https://www.hdfcfund.com/explore/mutual-funds/{slug}/{plan}"
    try:
        response = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
        response.raise_for_status()
    except requests.RequestException:
        return {"source_url": url, "source_name": "HDFC Mutual Fund" if is_hdfc else "Motilal Oswal AMC", "unavailable": True}
    raw = response.text
    soup = BeautifulSoup(raw, "html.parser")
    text = soup.get_text(" ", strip=True)

    def value(pattern, source=raw):
        found = re.search(pattern, source, re.I | re.S)
        return found.group(1).strip() if found else None

    if is_hdfc:
        # HDFC publishes scheme facts in the public scheme page and its embedded
        # portfolio data. Verify the page belongs to this scheme before parsing.
        expected = re.sub(r"[^a-z0-9]", "", base_name.lower())
        page_text = re.sub(r"[^a-z0-9]", "", text.lower())
        if expected not in page_text:
            return {"source_url": url, "source_name": "HDFC Mutual Fund", "unavailable": True}
        inception_raw = value(r"Inception Date\s+(\d{2}/\d{2}/\d{4})", text)
        try:
            inception = dt.datetime.strptime(inception_raw, "%d/%m/%Y").date().isoformat() if inception_raw else None
        except ValueError:
            inception = None
        ter = value(r"TER\s+TER\s+.*?Disclaimer.*?\s([\d.]+)\s+Lock in", text)
        aum_match = re.search(r"AUM\s*\((\d{2}/\d{2}/\d{4})\).*?₹\s*([\d,.]+)\s*Cr", text, re.I | re.S)
        lock_in = value(r"Lock in\s+(NA|[\w\s-]{1,35}?)\s+(?:NAV|Benchmark)", text)
        risk = value(r"The Risk of the s(?:c)?heme is\s+([A-Za-z -]{2,40})", text)
        manager_block = value(r"Fund Managers(.*?)Portfolio Allocation & Top Holdings", text)
        managers = []
        for manager in re.findall(r"(?:Mr\.?|Ms\.?|Mrs\.?)\s+(.+?)(?=\s+(?:Fund Manager|Dealer|Chief Investment|Equity Analyst|OVERSEAS|Mr\.?\s|Ms\.?\s|Mrs\.?\s|$))", manager_block or ""):
            clean = manager.strip(" .")
            if clean and not any(item["name"] == clean for item in managers):
                managers.append({"name": clean, "role": "Fund Manager"})
        exit_load = value(r"Exit Load\s+(.*?)(?:Product Labelling)", text)
        date_on = value(r"Portfolio Allocation & Top Holdings\s+As on\s+(\d{1,2}\s+[A-Za-z]{3}\s+\d{4})", text)
        holdings, sectors, portfolio_date = _hdfc_official_portfolio(base_name)
        date_on = portfolio_date or date_on
        return {
            "source_url": url, "source_name": "HDFC Mutual Fund", "as_of": date_on,
            "fund_age_start": inception, "aum_crore": aum_match.group(2) if aum_match else None,
            "aum_date": aum_match.group(1) if aum_match else None,
            "expense_ratio": (ter + "%") if ter and "%" not in ter else ter,
            "exit_load": " ".join((exit_load or "").split())[:500] or None,
            "lock_in": lock_in, "managers": managers, "holdings": holdings, "sectors": sectors,
            "riskometer": risk, "tax_implications": None,
        }

    def disclosure_rows(key, pairs):
        match = re.search(r"(?:[\"']" + key + r"[\"']|\b" + key + r")\s*:\s*\[(.*?)\]", raw, re.I | re.S)
        if not match:
            return []
        output = []
        for block in re.findall(r"\{(.*?)\}", match.group(1), re.S):
            row = {}
            for source_key, target_key in pairs:
                field = value(r"[\"']?" + source_key + r"[\"']?\s*:\s*[\"']([^\"']+)", block)
                if field:
                    row[target_key] = field
            if row:
                output.append(row)
        return output

    managers = disclosure_rows("fundManager", [("fundManagerName", "name"), ("designation", "role")])
    holdings = disclosure_rows("holdings", [("nameOfSecurity", "name"), ("percentToNAV", "allocation")])
    sectors = disclosure_rows("sector", [("sector", "name"), ("percentage", "allocation")])
    if is_motilal:
        # The AMC page exposes the same official disclosure both as embedded
        # component data and as rendered label/value text. Keep a text fallback
        # because its website has changed the embedded JSON formatting before.
        if not managers:
            manager_names = re.findall(r"fundManagerName\s*:\s*(.+?)\s+(?=type\s*:|dropdownField\s*:|cfReferencePath\s*:)", text, re.I)
            managers = [{"name": re.sub(r"^Mr\.\s*", "", item).strip(), "role": "Fund Manager"} for item in manager_names if item.strip()]
        if not holdings:
            holdings = [
                {"name": name.strip(), "allocation": allocation.strip()}
                for name, allocation in re.findall(r"nameOfSecurity\s*:\s*(.+?)\s+percentToNAV\s*:\s*([\d,.]+%?)", text, re.I)
                if name.strip() and allocation.strip()
            ]
        if not sectors:
            sectors = [
                {"name": name.strip(), "allocation": allocation.strip()}
                for name, allocation in re.findall(r"sector\s*:\s*(.+?)\s+percentage\s*:\s*([\d,.]+%?)", text, re.I)
                if name.strip() and allocation.strip()
            ]
    portfolio_date = value(r"Date As On\s*: ?\s*([0-9]{1,2}\s+[A-Za-z]{3}\s+20[0-9]{2})", text)
    if is_motilal and not holdings:
        workbook_holdings, workbook_sectors, workbook_date = _motilal_official_portfolio(raw, url)
        holdings = workbook_holdings
        sectors = sectors or workbook_sectors
        portfolio_date = portfolio_date or workbook_date
    facts = {
        "source_url": url, "source_name": "Motilal Oswal AMC",
        "as_of": portfolio_date,
        "fund_age_start": value(r"[\"']?dateOfAllotment[\"']?\s*:\s*[\"']([^\"']+)") or value(r"[\"']?nfoStartDate[\"']?\s*:\s*[\"']([^\"']+)") or value(r"Date of inception\s*:?\s*(\d{1,2}-[A-Za-z]{3}-\d{4})", text),
        "aum_crore": value(r"[\"']?latestAum[\"']?\s*:\s*[\"']?([\d,.]+)") or value(r"Latest AUM\s+([\d,.]+)", text),
        "aum_date": value(r"[\"']?latestAumAsOnDt[\"']?\s*:\s*[\"']([^\"']+)") or value(r"[\"']?latestAumAsOnDt[\"']?\s*:\s*([\dT:-]+)") or value(r"Latest AUM\s+[\d,.]+\s+(\d{4}-\d{2}-\d{2})", text),
        "expense_ratio": (
            value(r"[\"']?expenseRatioDirect[\"']?\s*:\s*[\"']?([\d.]+%?)") if "direct" in str(scheme.get("plan") or "").lower()
            else value(r"[\"']?expenseRatioRegular[\"']?\s*:\s*[\"']?([\d.]+%?)")
        ) or value(r"Total Expense Ratio\s*([\d.]+%)", text),
        "exit_load": value(r"[\"']?exitLoad[\"']?\s*:\s*[\"']([^\"']+)") or value(r"Exit Load Policy\s+(.{1,260}?)(?:Entry Load|Minimum Application Amount)", text),
        "lock_in": value(r"[\"']?lockInPeriod[\"']?\s*:\s*[\"']([^\"']+)") or value(r"Lock.?in(?: Period)?\s*:?\s*([^\n]{1,80})", text),
        "managers": managers, "holdings": holdings,
        "sectors": sectors,
        "riskometer": value(r"[\"']?riskType[\"']?\s*:\s*[\"']([^\"']+)") or value(r"Scheme Risk.?O.?Meter\s+([A-Za-z ]{3,30})", text),
        # Tax treatment changes with regulations and investor circumstances; do not
        # display an unsourced generic tax statement as if it were scheme-specific.
        "tax_implications": None,
    }


def _fetch_angel_one_fund_facts(scheme):
    """Read SBI scheme facts from the matching Angel One fund detail page."""
    scheme_name = str(scheme.get("name") or "").strip()
    page_name = re.sub(r"\s*-\s*(?:direct|regular)\s+plan\b.*$", "", scheme_name, flags=re.I).strip()
    page_name = re.sub(r"\s*\((?:direct|regular).*?\)\s*$", "", page_name, flags=re.I).strip()
    plan = str(scheme.get("plan") or "").lower()
    if plan not in ("direct", "regular"):
        plan_match = re.search(r"\b(direct|regular)\s+plan\b", scheme_name, re.I)
        plan = plan_match.group(1).lower() if plan_match else "direct"
    option = str(scheme.get("option") or "").lower()
    if not option:
        option = scheme_name.lower()
    if "reinvest" in option:
        option_slug = "idcw-reinvestment"
    elif "payout" in option or "dividend" in option:
        option_slug = "idcw-payout"
    else:
        option_slug = "growth"
    base_slug = re.sub(r"[^a-z0-9]+", "-", page_name.replace("&", " and ").lower()).strip("-")
    slug = base_slug + "-" + plan + "-plan-" + option_slug
    url = "https://www.angelone.in/mutual-funds/mf-schemes/" + slug
    try:
        response = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
        response.raise_for_status()
    except requests.RequestException:
        return {"source_url": url, "source_name": "Angel One", "unavailable": True}

    soup = BeautifulSoup(response.text, "html.parser")
    text = soup.get_text(" ", strip=True)
    expected = re.sub(r"[^a-z0-9]", "", page_name.lower())
    normalized = re.sub(r"[^a-z0-9]", "", text.lower())
    if not expected or expected not in normalized:
        return {"source_url": url, "source_name": "Angel One", "unavailable": True}

    def find(pattern, source=text):
        match = re.search(pattern, source, re.I | re.S)
        return " ".join(match.group(1).split()) if match else None

    age_match = re.search(r"Launched on\s+([A-Za-z]+\s+20\d{2})\s*\((\d+)\s*years?\)", text, re.I)
    aum = find(r"Asset Under Management\s+₹\s*([\d,.]+)\s*Cr")
    expense = find(r"Expense Ratio\s+([\d.]+%)\s*\(inclusive of GST\)") or find(r"Expense Ratio\s+([\d.]+%)")
    exit_load = find(r"Exit Load\s+(.+?)(?=Ratings|Tax Implications|SBI Mutual Fund Manager)")
    lock_in = find(r"Fund has\s+(.+?)\s+lock-in period")
    if lock_in:
        lock_in = "No lock-in period" if lock_in.strip().lower() == "no" else lock_in + " lock-in period"
    tax = None
    tax_match = re.search(
        r"Tax Implications\s+Withdrawal within 1 year:\s*(.+?)\s+Withdrawal after 1 year:\s*(.+?)\s+(?:SBI Mutual Fund Manager|Fund House Details)",
        text, re.I | re.S,
    )
    if tax_match:
        tax = "Withdrawal within 1 year: " + " ".join(tax_match.group(1).split()) + "; Withdrawal after 1 year: " + " ".join(tax_match.group(2).split())

    manager_block = find(r"SBI Mutual Fund Manager\s+(.+?)\s+Fund House Details") or ""
    managers = []
    manager_matches = re.findall(r"([A-Z][A-Za-z. ]{1,50}?)\s+Fund Manager since\s+([A-Za-z]+\s+20\d{2})", manager_block)
    for person, since in manager_matches:
        clean = re.sub(r"^(?:[A-Z]{1,3}\s+)+", "", person).strip()
        if clean:
            managers.append({"name": clean, "role": "Fund Manager since " + since})

    ratings = []
    for label, pattern in (
        ("Value Research", r"Value Research\s+(\d(?:\.\d+)?)"),
        ("Crisil", r"Crisil\s+(\d(?:\.\d+)?)"),
        ("Morningstar", r"Morning Star\s+(\d(?:\.\d+)?)"),
    ):
        rating = find(pattern)
        if rating:
            ratings.append({"name": label, "rating": rating})

    return {
        "source_url": url,
        "source_name": "Angel One",
        # The scheme page has a NAV date, but it is not a holdings disclosure
        # date. Do not reuse it as one in the portfolio section.
        "as_of": None,
        "fund_age_start": None,
        "fund_age_years": int(age_match.group(2)) if age_match else None,
        "fund_launch_label": age_match.group(1) if age_match else None,
        "aum_crore": aum,
        "aum_date": None,
        "expense_ratio": expense,
        "exit_load": exit_load,
        "lock_in": lock_in,
        "tax_implications": tax,
        "managers": managers,
        "holdings": [],
        "sectors": [],
        "holdings_notice": "Angel One’s scheme page does not publish stock or sector allocations.",
        "riskometer": find(r"Your principal will be at\s+([A-Za-z ]+Risk)"),
        "ratings": ratings,
    }


def _hdfc_official_portfolio(base_name):
    """Download and parse the latest matching HDFC monthly portfolio workbook."""
    listing_url = "https://www.hdfcfund.com/statutory-disclosure/portfolio/monthly-portfolio"
    try:
        listing = requests.get(listing_url, headers=HEADERS, timeout=TIMEOUT)
        listing.raise_for_status()
        soup = BeautifulSoup(listing.text, "html.parser")
        target = re.sub(r"[^a-z0-9]", "", base_name.lower())
        candidates = []
        for anchor in soup.find_all("a", href=True):
            label = anchor.get_text(" ", strip=True)
            normalized = re.sub(r"[^a-z0-9]", "", label.lower())
            href = urljoin(listing.url, anchor["href"])
            if target and target in normalized and href.lower().split("?")[0].endswith(".xlsx"):
                candidates.append((label, href))
        if not candidates:
            return [], [], None
        label, file_url = candidates[0]
        book_response = requests.get(file_url, headers=HEADERS, timeout=TIMEOUT)
        book_response.raise_for_status()
        book = load_workbook(BytesIO(book_response.content), read_only=True, data_only=True)
    except Exception:
        return [], [], None
    holdings = []
    sector_weights = {}

    def formatted_percent(value):
        raw_value = str(value or "").strip()
        if not raw_value:
            return ""
        if "%" in raw_value:
            return raw_value
        try:
            numeric = float(raw_value.replace(",", ""))
        except ValueError:
            return raw_value
        if 0 < numeric < 1:
            numeric *= 100
        return f"{numeric:.2f}%"

    for sheet in book.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header_row = None
        header_index = 0
        for idx, row in enumerate(rows):
            labels = [re.sub(r"[^a-z0-9]", "", str(cell or "").lower()) for cell in row]
            name_idx = next((i for i, item in enumerate(labels) if any(key in item for key in ("nameofsecurity", "securityname", "companyname", "issuername", "instrumentname"))), None)
            pct_idx = next((i for i, item in enumerate(labels) if any(key in item for key in ("percenttonav", "percentofnav", "percenttonetassets", "netassetspercent", "navpercent", "tonetassets", "tonav"))), None)
            if name_idx is not None and pct_idx is not None:
                header_row, header_index = row, idx
                sector_idx = next((i for i, item in enumerate(labels) if "sector" in item or "industry" in item), None)
                break
        if header_row is None:
            continue
        for row in sheet.iter_rows(min_row=header_index + 2, values_only=True):
            if len(row) <= max(name_idx, pct_idx):
                continue
            security = str(row[name_idx] or "").strip()
            if not security or security.lower() in ("total", "grand total", "subtotal"):
                continue
            allocation = formatted_percent(row[pct_idx])
            if not allocation:
                continue
            holdings.append({"name": security, "allocation": allocation if "%" in allocation else allocation + "%"})
            if sector_idx is not None and len(row) > sector_idx and row[sector_idx]:
                try:
                    weight = float(allocation.replace("%", "").replace(",", ""))
                    sector = str(row[sector_idx]).strip()
                    sector_weights[sector] = sector_weights.get(sector, 0) + weight
                except (TypeError, ValueError):
                    pass
        if holdings:
            break
    sectors = [{"name": name, "allocation": f"{weight:.2f}%"} for name, weight in sorted(sector_weights.items(), key=lambda pair: pair[1], reverse=True)]
    date_match = re.search(r"(\d{1,2}\s+[A-Za-z]+\s+\d{4})", label)
    return holdings, sectors, date_match.group(1) if date_match else None


def _motilal_official_portfolio(raw, page_url):
    """Parse Motilal Oswal's linked month-end Excel portfolio disclosure."""
    match = re.search(r"portfolioUrl\s*:\s*[\"']?([^\"'\s]+\.xlsx)", raw, re.I)
    if not match:
        return [], [], None
    file_url = urljoin(page_url, match.group(1).replace("\\/", "/"))
    try:
        response = requests.get(file_url, headers=HEADERS, timeout=TIMEOUT)
        response.raise_for_status()
        book = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    except Exception:
        return [], [], None
    holdings = []
    sectors = []
    for sheet in book.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        for row_index, row in enumerate(rows):
            values = [str(value or "").strip() for value in row]
            lowered = [re.sub(r"[^a-z0-9]", "", value.lower()) for value in values]
            security_idx = next((i for i, value in enumerate(lowered) if value in ("nameofsecurity", "securityname", "name")), None)
            weight_idx = next((i for i, value in enumerate(lowered) if value in ("percenttonav", "percentofnav", "percentage")), None)
            if security_idx is not None and weight_idx is not None:
                for item in rows[row_index + 1:]:
                    if len(item) <= max(security_idx, weight_idx) or not item[security_idx] or not item[weight_idx]:
                        continue
                    holdings.append({"name": str(item[security_idx]).strip(), "allocation": str(item[weight_idx]).strip()})
                break
        if holdings:
            break
    date_match = re.search(r"Date As On\s*:?\s*([0-9]{1,2}\s+[A-Za-z]{3}\s+20[0-9]{2})", raw, re.I)
    return holdings, sectors, date_match.group(1) if date_match else None


def _follow_history_frames(client, response, headers, depth=0, timeout=None):
    """Follow AMFI's legacy frameset wrapper to its actual history report."""
    if depth >= 3:
        return {}
    soup = BeautifulSoup(_decode(response), "html.parser")
    for frame in soup.find_all(["frame", "iframe"]):
        source = str(frame.get("src") or "").strip()
        if not source or source.startswith(("#", "about:", "javascript:")):
            continue
        target = urljoin(response.url, source)
        host = (urlparse(target).hostname or "").lower()
        if host != "amfiindia.com" and not host.endswith(".amfiindia.com"):
            continue
        frame_headers = dict(headers)
        frame_headers["Referer"] = response.url
        try:
            child = client.get(target, headers=frame_headers, timeout=timeout or TIMEOUT)
            child.raise_for_status()
        except requests.RequestException:
            continue
        rows = _parse_history_rows(_decode(child))
        if rows:
            return rows
        rows = _follow_history_frames(client, child, frame_headers, depth + 1, timeout)
        if rows:
            return rows
    return {}


def _history_label(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _fetch_history_from_amfi_form(start_date, end_date, headers, amc_name=None, scheme_type=None, timeout=None, fast=False):
    """Use the current AMFI NAV History form if its download route has changed."""
    try:
        client = requests.Session()
        page = client.get(AMFI_NAV_PAGE, headers=headers, timeout=timeout or TIMEOUT)
        page.raise_for_status()
    except requests.RequestException:
        return None

    soup = BeautifulSoup(_decode(page), "html.parser")
    forms = []
    for form in soup.find_all("form"):
        heading = form.find_previous(["h2", "h3", "h4"])
        heading_text = heading.get_text(" ", strip=True).lower() if heading else ""
        context = (heading_text + " " + form.get_text(" ", strip=True)).lower()
        if "history" not in context or "old version" in heading_text or "old version" in context:
            continue
        controls = form.find_all(["input", "select"])
        names = " ".join(str(control.get("name") or control.get("id") or "").lower() for control in controls)
        if any(key in names for key in ("from", "frm", "start")) and any(key in names for key in ("to", "end")):
            forms.append(form)

    for form in forms:
        base_values = {}
        from_name = None
        to_name = None
        found_amc = False
        found_type = False
        select_controls = form.find_all("select")
        amc_controls = [control for control in select_controls if any(token in str(control.get("name") or control.get("id") or "").lower() for token in ("amc", "fund", "mutual"))]
        type_controls = [control for control in select_controls if "type" in str(control.get("name") or control.get("id") or "").lower()]
        for control in form.find_all("input"):
            name = control.get("name") or control.get("id")
            if not name:
                continue
            kind = str(control.get("type") or "text").lower()
            if kind in ("submit", "button", "image", "reset"):
                continue
            key = name.lower()
            if any(part in key for part in ("from", "frmdt", "frmdate", "startdate")):
                from_name = name
            elif any(part in key for part in ("to", "todt", "todate", "enddate")):
                to_name = name
            else:
                base_values[name] = control.get("value", "")
        for select in select_controls:
            name = select.get("name") or select.get("id")
            if not name:
                continue
            selector_name = name.lower()
            is_amc_selector = any(token in selector_name for token in ("amc", "fund", "mutual"))
            is_type_selector = "type" in selector_name
            options = select.find_all("option")
            choice = next((option for option in options if "all" in option.get_text(" ", strip=True).lower()), None)
            if choice is not None:
                if is_amc_selector or not is_type_selector:
                    found_amc = True
                if is_type_selector or not is_amc_selector:
                    found_type = True
            option_values = [(option, _history_label(option.get_text(" ", strip=True))) for option in options]
            if choice is None and amc_name and (is_amc_selector or not is_type_selector):
                wanted_amc = _history_label(amc_name.replace("Mutual Fund", ""))
                choice = next((option for option, label in option_values if wanted_amc and wanted_amc in label), None)
                found_amc = found_amc or choice is not None
            if choice is None and scheme_type and (is_type_selector or not is_amc_selector):
                wanted_type = _history_label(scheme_type)
                choice = next((option for option, label in option_values if wanted_type and (wanted_type in label or label in wanted_type)), None)
                found_type = found_type or choice is not None
            if choice is not None:
                base_values[name] = choice.get("value", "")
        if amc_controls and not found_amc:
            # Do not assume the first AMC in AMFI's selector is the requested fund.
            continue
        if type_controls and not found_type:
            continue
        if select_controls and not amc_controls and not type_controls and not (found_amc and found_type):
            continue
        if not from_name or not to_name:
            continue
        if amc_name and not found_amc:
            continue
        if scheme_type and not found_type:
            continue

        action = urljoin(page.url, form.get("action") or page.url)
        method = str(form.get("method") or "get").lower()
        date_formats = ("%d-%b-%Y",) if fast else ("%d-%b-%Y", "%Y-%m-%d")
        for date_format in date_formats:
            values = dict(base_values)
            values[from_name] = start_date.strftime(date_format)
            values[to_name] = end_date.strftime(date_format)
            submit = next((control for control in form.find_all(["button", "input"]) if "go" in (control.get_text(" ", strip=True) + " " + str(control.get("value") or "")).lower()), None)
            if submit and submit.get("name"):
                values[submit["name"]] = submit.get("value") or submit.get_text(" ", strip=True)
            try:
                if method == "post":
                    response = client.post(action, data=values, headers=headers, timeout=timeout or TIMEOUT)
                else:
                    response = client.get(action, params=values, headers=headers, timeout=timeout or TIMEOUT)
                response.raise_for_status()
                rows = _parse_history_rows(_decode(response))
                if not rows:
                    rows = _follow_history_frames(client, response, headers, timeout=timeout)
                if rows:
                    return rows
            except requests.RequestException:
                continue
    return None


def fetch_history(start_date, end_date, amc_name=None, scheme_type=None, form_only=False, timeout=None, fast_form=False):
    """Fetch up to 90 days of AMFI history for all schemes; caller filters codes."""
    if (end_date - start_date).days > 89:
        raise AMFIError("AMFI history requests are limited to 90 days.")
    start_text = start_date.strftime("%d-%b-%Y")
    end_text = end_date.strftime("%d-%b-%Y")
    urls = [
        "https://portal.amfiindia.com/DownloadNAVHistoryReport_Po.aspx?"
        f"frmdt={start_text}&todt={end_text}",
        "https://portal.amfiindia.com/DownloadNAVHistoryReport_Po.aspx?"
        f"tp=1&frmdt={start_text}&todt={end_text}",
        "https://portal.amfiindia.com/DownloadNAVHistoryReport_Po.aspx?"
        f"tp=1&frmdt={start_text}&todt={end_text}&mf=all",
        "https://portal.amfiindia.com/NavHistoryReport_Rpt_Po.aspx?"
        f"frmdate={start_text}&rpt=0",
        "https://portal.amfiindia.com/NavHistoryReport_Rpt_Po.aspx?"
        f"frmdate={start_text}&mf=all&rpt=dn",
        "https://portal.amfiindia.com/NavHistoryReport_Rpt_Po.aspx?"
        f"frmdate={start_text}&rpt=dn",
    ]
    errors = []
    request_timeout = timeout or TIMEOUT
    request_headers = dict(HEADERS)
    request_headers.update({
        "Referer": AMFI_NAV_PAGE,
        "Origin": "https://www.amfiindia.com",
    })
    if not form_only:
        client = requests.Session()
        for url in urls:
            try:
                response = client.get(url, headers=request_headers, timeout=request_timeout)
                response.raise_for_status()
                text = _decode(response)
                rows = _parse_history_rows(text)
                if not rows:
                    rows = _follow_history_frames(client, response, request_headers, timeout=request_timeout)
                if rows:
                    return rows
                content_type = response.headers.get("Content-Type", "unknown content type")
                sample = " ".join(text[:140].split())
                errors.append(f"AMFI returned {response.status_code} ({content_type}) without NAV rows: {sample}")
            except requests.RequestException as exc:
                errors.append(str(exc))
    form_rows = _fetch_history_from_amfi_form(
        start_date, end_date, request_headers, amc_name, scheme_type,
        timeout=request_timeout, fast=fast_form,
    )
    if form_rows:
        return form_rows
    raise AMFIError("Could not retrieve historical NAVs from AMFI. " + " ".join(errors[-2:]))
